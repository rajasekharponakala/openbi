"""Dataset refresh: run every table's query and store the result.

Import-mode tables are materialised as DuckDB tables in the model's own
database file (the columnar, compressed analogue of VertiPaq). DirectQuery
tables become views evaluated at query time. A refresh builds a new file
and atomically swaps it in, so readers never see a half-refreshed model.
"""

from __future__ import annotations

import csv
import json
import os
import tempfile
import time
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import duckdb

from ..dax.compiler import RowCompiler
from ..dax.parser import parse
from ..engine.graph import ModelGraph
from ..model.schema import DUCKDB_TYPES, SemanticModel, Source, Step, Table
from .sql import ident, literal, select_query
from .steps import compile_steps, step_dependencies

ATTACH_TYPES = {"postgres": "POSTGRES", "postgresql": "POSTGRES", "mysql": "MYSQL", "sqlite": "SQLITE", "duckdb": None}


class RefreshError(RuntimeError):
    pass


@dataclass
class TableStats:
    name: str
    rows: int | None
    mode: str
    seconds: float


@dataclass
class RefreshResult:
    started_at: str
    seconds: float
    tables: list[TableStats] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {
            "startedAt": self.started_at,
            "seconds": round(self.seconds, 3),
            "tables": [t.__dict__ for t in self.tables],
        }


def _resolve_path(path: str, base_dir: Path, allowed_root: Path | None) -> Path:
    p = Path(path)
    if not p.is_absolute():
        p = base_dir / p
    p = p.resolve()
    if allowed_root is not None:
        root = allowed_root.resolve()
        if not p.is_relative_to(root):
            raise RefreshError(f"path '{path}' is outside the allowed data directory")
    if "*" not in p.name and not p.exists():
        raise RefreshError(f"file not found: {p}")
    return p


def _attach_statement(alias: str, connection: str) -> str:
    kind, _, rest = connection.partition(":")
    kind = kind.lower()
    if kind not in ATTACH_TYPES:
        raise RefreshError(f"unsupported database connection '{kind}' (use postgres:, mysql:, sqlite: or duckdb:)")
    t = ATTACH_TYPES[kind]
    opts = f"(TYPE {t}, READ_ONLY)" if t else "(READ_ONLY)"
    return f"ATTACH {literal(rest.lstrip('/') if kind in ('sqlite', 'duckdb') else rest)} AS {ident(alias)} {opts}"


def _excel_to_csv(path: Path, sheet: str | None, header: bool, out_dir: Path) -> Path:
    from openpyxl import load_workbook  # thin adapter: openpyxl reads, DuckDB parses CSV

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    out = out_dir / f"{path.stem}-{ws.title}.csv"
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        for row in ws.iter_rows(values_only=True):
            w.writerow(["" if v is None else (v.isoformat() if hasattr(v, "isoformat") else v) for v in row])
    wb.close()
    return out


def sandbox(con: duckdb.DuckDBPyConnection, directories: list[Path]) -> None:
    """Restrict DuckDB file access to ``directories`` (DuckDB's own sandbox)."""
    dirs = ", ".join(literal(str(d.resolve()) + os.sep) for d in directories)
    con.execute(f"SET allowed_directories = [{dirs}]")
    con.execute("SET enable_external_access = false")
    con.execute("SET lock_configuration = true")


class SourceCompiler:
    """Turns a :class:`Source` into a FROM-able SQL expression."""

    def __init__(self, con: duckdb.DuckDBPyConnection, base_dir: Path, allowed_root: Path | None, scratch: Path):
        self.con = con
        self.base_dir = base_dir
        self.allowed_root = allowed_root
        self.scratch = scratch
        self.attachments: list[str] = []

    def sql(self, table: str, src: Source) -> str:
        if src.kind in ("csv", "parquet", "json", "excel"):
            path = _resolve_path(src.path or "", self.base_dir, self.allowed_root)
            if src.kind == "excel":
                path = _excel_to_csv(path, src.sheet, src.header, self.scratch)
                return f"SELECT * FROM read_csv_auto({literal(str(path))}, header = {literal(src.header)})"
            if src.kind == "csv":
                opts = f", header = {literal(src.header)}"
                if src.delimiter:
                    opts += f", delim = {literal(src.delimiter)}"
                return f"SELECT * FROM read_csv_auto({literal(str(path))}{opts})"
            if src.kind == "parquet":
                return f"SELECT * FROM read_parquet({literal(str(path))})"
            return f"SELECT * FROM read_json_auto({literal(str(path))})"
        if src.kind == "inline":
            rows = src.rows or []
            if not rows:
                raise RefreshError(f"inline source of '{table}' has no rows")
            cols = list(dict.fromkeys(k for r in rows for k in r))
            values = ", ".join("(" + ", ".join(literal(r.get(c)) for c in cols) + ")" for r in rows)
            return f"SELECT * FROM (VALUES {values}) AS _v({', '.join(ident(c) for c in cols)})"
        if src.kind == "sql":
            return select_query(src.query or "")
        if src.kind == "database":
            alias = f"src_{len(self.attachments)}"
            stmt = _attach_statement(alias, src.connection or "")
            self.con.execute(stmt)
            self.attachments.append(stmt)
            if src.query:
                return select_query(src.query)
            schema, _, name = (src.table or "").rpartition(".")
            target = f"{ident(alias)}.{ident(schema)}.{ident(name)}" if schema else f"{ident(alias)}.{ident(name)}"
            return f"SELECT * FROM {target}"
        raise RefreshError(f"unknown source kind {src.kind}")  # pragma: no cover


def _ordered(model: SemanticModel) -> list[Table]:
    by_name = {t.name.lower(): t for t in model.tables}
    order: list[Table] = []
    state: dict[str, int] = {}

    def visit(t: Table) -> None:
        key = t.name.lower()
        if state.get(key) == 2:
            return
        if state.get(key) == 1:
            raise RefreshError(f"circular query dependency involving '{t.name}'")
        state[key] = 1
        for dep in step_dependencies(t.steps):
            d = by_name.get(dep.lower())
            if d is None:
                raise RefreshError(f"table '{t.name}' references unknown query '{dep}'")
            visit(d)
        state[key] = 2
        order.append(t)

    for t in model.tables:
        visit(t)
    return order


def table_query(table: Table, source_sql: str, upto: int | None = None) -> str:
    steps = table.steps if upto is None else table.steps[:upto]
    typed = {c.name: c.data_type.value for c in table.columns if c.data_type and not c.expression}
    if typed and upto is None:
        steps = steps + [Step(op="changeType", types=typed)]
    return compile_steps(source_sql, steps, lambda n: ident(n))


def _add_calculated_columns(con: duckdb.DuckDBPyConnection, model: SemanticModel, table: Table) -> None:
    graph = ModelGraph(model)
    obj = "TABLE" if table.mode == "import" else "VIEW"
    for c in table.columns:
        if not c.expression:
            continue
        rc = RowCompiler(model, graph, table.name)
        expr = rc.expr(parse(c.expression))
        joins = ""
        reach = graph.reachable(table.name)
        seen = {graph.canonical(table.name)}
        for t in rc.tables:
            for e in reach[t]:
                if e.dst not in seen:
                    seen.add(e.dst)
                    joins += f" LEFT JOIN {ident(e.dst)} ON {ident(e.src)}.{ident(e.src_col)} = {ident(e.dst)}.{ident(e.dst_col)}"
        if c.data_type:
            expr = f"TRY_CAST({expr} AS {DUCKDB_TYPES[c.data_type]})"
        name = ident(table.name)
        tmp = ident(f"__calc_{table.name}")
        con.execute(f"ALTER {obj} {name} RENAME TO {tmp}")
        # keep the base alias equal to the table name so the expression resolves
        con.execute(f"CREATE {obj} {name} AS SELECT {name}.*, {expr} AS {ident(c.name)} FROM {tmp} AS {name}{joins}")
        if obj == "TABLE":
            con.execute(f"DROP TABLE {tmp}")


def refresh_model(
    model: SemanticModel,
    db_path: str | Path,
    base_dir: str | Path = ".",
    allowed_root: str | Path | None = None,
) -> RefreshResult:
    db_path = Path(db_path)
    db_path.parent.mkdir(parents=True, exist_ok=True)
    started = datetime.now(UTC).isoformat()
    t0 = time.perf_counter()
    result = RefreshResult(started, 0.0)
    fd, tmp_name = tempfile.mkstemp(prefix=db_path.name + ".", suffix=".refresh", dir=db_path.parent)
    os.close(fd)
    os.unlink(tmp_name)  # duckdb must create the file itself
    try:
        with tempfile.TemporaryDirectory() as scratch:
            con = duckdb.connect(tmp_name)
            try:
                if allowed_root is not None and not _uses_databases(model):
                    sandbox(con, [Path(allowed_root), Path(scratch)])
                sc = SourceCompiler(con, Path(base_dir), Path(allowed_root) if allowed_root else None, Path(scratch))
                for table in _ordered(model):
                    ts = time.perf_counter()
                    try:
                        sql = table_query(table, sc.sql(table.name, table.source))
                        if table.mode == "import":
                            con.execute(f"CREATE TABLE {ident(table.name)} AS {sql}")
                        elif table.source.kind == "excel":
                            raise RefreshError(f"DirectQuery is not available for {table.source.kind} sources")
                        else:
                            con.execute(f"CREATE VIEW {ident(table.name)} AS {sql}")
                    except duckdb.Error as e:
                        raise RefreshError(f"table '{table.name}': {e}") from e
                    result.tables.append(TableStats(table.name, None, table.mode, time.perf_counter() - ts))
                for table in model.tables:
                    try:
                        _add_calculated_columns(con, model, table)
                    except (duckdb.Error, ValueError) as e:
                        raise RefreshError(f"calculated column in '{table.name}': {e}") from e
                for st in result.tables:
                    if st.mode == "import":
                        st.rows = con.execute(f"SELECT COUNT(*) FROM {ident(st.name)}").fetchone()[0]
                result.seconds = time.perf_counter() - t0
                con.execute("CREATE TABLE __openbi_meta (key VARCHAR PRIMARY KEY, value VARCHAR)")
                con.execute(
                    "INSERT INTO __openbi_meta VALUES (?, ?), (?, ?)",
                    ["refresh", json.dumps(result.as_dict()), "attachments", json.dumps(sc.attachments)],
                )
                con.execute("CHECKPOINT")
            finally:
                con.close()
        os.replace(tmp_name, db_path)
    except BaseException:
        for p in (tmp_name, tmp_name + ".wal"):
            if os.path.exists(p):
                os.unlink(p)
        raise
    return result


def preview_table(
    model: SemanticModel,
    table_name: str,
    base_dir: str | Path = ".",
    allowed_root: str | Path | None = None,
    upto: int | None = None,
    limit: int = 100,
) -> dict[str, Any]:
    """Evaluate a table's query up to step ``upto`` (Power Query editor preview)."""
    target = model.table(table_name)
    if target is None:
        raise RefreshError(f"unknown table '{table_name}'")
    con = duckdb.connect(":memory:")
    try:
        with tempfile.TemporaryDirectory() as scratch:
            if allowed_root is not None and not _uses_databases(model):
                sandbox(con, [Path(allowed_root), Path(scratch)])
            sc = SourceCompiler(con, Path(base_dir), Path(allowed_root) if allowed_root else None, Path(scratch))
            for table in _ordered(model):
                if table is target:
                    break
                if table.name.lower() in {d.lower() for d in _all_deps(model, target)}:
                    con.execute(
                        f"CREATE VIEW {ident(table.name)} AS {table_query(table, sc.sql(table.name, table.source))}"
                    )
            sql = table_query(target, sc.sql(target.name, target.source), upto)
            try:
                cur = con.execute(f"SELECT * FROM ({sql}) LIMIT {int(limit)}")
            except duckdb.Error as e:
                raise RefreshError(str(e)) from e
            rows = cur.fetchall()
            from ..engine.query import _jsonable

            return {
                "columns": [{"name": d[0], "type": str(d[1])} for d in cur.description],
                "rows": [[_jsonable(v) for v in r] for r in rows],
                "sql": sql,
            }
    finally:
        con.close()


def _uses_databases(model: SemanticModel) -> bool:
    return any(t.source.kind == "database" for t in model.tables)


def _all_deps(model: SemanticModel, table: Table) -> set[str]:
    out: set[str] = set()
    stack = [table]
    while stack:
        t = stack.pop()
        for d in step_dependencies(t.steps):
            if d.lower() not in out:
                out.add(d.lower())
                dt = model.table(d)
                if dt:
                    stack.append(dt)
    return out


def open_model_db(db_path: str | Path) -> duckdb.DuckDBPyConnection:
    """Open a refreshed model read-only, re-attaching DirectQuery sources."""
    con = duckdb.connect(str(db_path), read_only=True)
    try:
        row = con.execute("SELECT value FROM __openbi_meta WHERE key = 'attachments'").fetchone()
    except duckdb.CatalogException:
        row = None
    attachments = json.loads(row[0]) if row else []
    for stmt in attachments:
        con.execute(stmt)
    if not attachments:
        # queries only read the model's own tables
        con.execute("SET enable_external_access = false")
        con.execute("SET lock_configuration = true")
    return con


def describe(con: duckdb.DuckDBPyConnection, model: SemanticModel) -> dict[str, list[dict[str, str]]]:
    """Physical columns and types per model table."""
    out: dict[str, list[dict[str, str]]] = {}
    for t in model.tables:
        try:
            rows = con.execute(f"DESCRIBE {ident(t.name)}").fetchall()
        except duckdb.Error:
            rows = []
        out[t.name] = [{"name": r[0], "type": r[1]} for r in rows]
    return out
